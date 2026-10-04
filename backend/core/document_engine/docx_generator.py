"""docx_generator — 规则驱动的合规 Word 生成（V4.5）

策略:
1. content_md 解析为结构化块（标题/段落/列表/表格）
2. FormatRules 应用到样式层（Normal/Heading N/页面边距）——样式级设置优先于逐段设置，
   生成的文档天然合规且用户后续在 Word 里打字也自动带格式
3. 生成后调 audit_docx 自检，返回合规率（闭环质量保证）

中文字体: 同时设置 ascii 与 eastAsia（python-docx 需手动操作 rFonts）
"""
import re
import logging
from pathlib import Path
from typing import Dict, Any, List, Optional

from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn

from core.document_engine.format_rules import FormatRules
from core.document_engine.docx_audit import audit_docx

logger = logging.getLogger(__name__)


class GenerateError(Exception):
    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.status_code = status_code


def _set_style_font(style: Any, name: Optional[str], size_pt: Optional[float],
                    bold: Optional[bool] = None) -> None:
    """样式字体设置（含中文 eastAsia）"""
    if name:
        style.font.name = name
        # 中文字体需写入 eastAsia
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(rfonts)
        rfonts.set(qn("w:eastAsia"), name)
    if size_pt is not None:
        style.font.size = Pt(size_pt)
    if bold is not None:
        style.font.bold = bold


def _parse_md_blocks(content_md: str) -> List[Dict[str, Any]]:
    """Markdown → 结构化块 [{type: heading/para/list_item/table_row...}]"""
    blocks: List[Dict[str, Any]] = []
    table_buffer: List[List[str]] = []

    def flush_table():
        nonlocal table_buffer
        if table_buffer:
            # 第一行是表头，第二行是分隔线（跳过）
            rows = table_buffer
            if len(rows) >= 2 and all(set(c) <= {"-", ":", " ", ""} for c in rows[1]):
                rows = [rows[0]] + rows[2:]
            blocks.append({"type": "table", "rows": rows})
            table_buffer = []

    for line in content_md.splitlines():
        stripped = line.strip()
        # 表格行
        if stripped.startswith("|") and stripped.endswith("|"):
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            table_buffer.append(cells)
            continue
        flush_table()

        if not stripped:
            continue
        m = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if m:
            blocks.append({"type": "heading", "level": len(m.group(1)), "text": m.group(2).strip()})
            continue
        m = re.match(r"^[-*]\s+(.+)$", stripped)
        if m:
            blocks.append({"type": "list_item", "text": m.group(1).strip()})
            continue
        m = re.match(r"^(\d+)[.、)]\s+(.+)$", stripped)
        if m:
            blocks.append({"type": "list_item", "text": m.group(2).strip(), "ordered": True})
            continue
        blocks.append({"type": "para", "text": stripped})

    flush_table()
    return blocks


def generate_docx(content_md: str, rules: FormatRules, output_path: Path,
                  title: Optional[str] = None) -> Dict[str, Any]:
    """生成合规 Word 并自检

    Args:
        content_md: Markdown 内容
        rules: 格式规则
        output_path: 输出 docx 路径（workspace 内，由调用方校验）
        title: 可选文档主标题（作为 H0 居中，不占标题层级）
    """
    if not content_md or not content_md.strip():
        raise GenerateError("content_md 不能为空", 400)

    doc = Document()

    # ── 1. 页面设置 ──
    if rules.margins:
        for sec in doc.sections:
            if rules.margins.top_cm is not None:
                sec.top_margin = Cm(rules.margins.top_cm)
            if rules.margins.bottom_cm is not None:
                sec.bottom_margin = Cm(rules.margins.bottom_cm)
            if rules.margins.left_cm is not None:
                sec.left_margin = Cm(rules.margins.left_cm)
            if rules.margins.right_cm is not None:
                sec.right_margin = Cm(rules.margins.right_cm)

    # ── 2. 样式层应用规则（关键: 让"后续打字也合规"）──
    normal = doc.styles["Normal"]
    if rules.body_font:
        _set_style_font(normal, rules.body_font.name, rules.body_font.size_pt)
    if rules.body_line_spacing is not None:
        normal.paragraph_format.line_spacing = rules.body_line_spacing

    heading_levels = {int(k): v for k, v in (rules.heading_styles or {}).items()}
    for lvl in range(1, 7):
        style_name = f"Heading {lvl}"
        try:
            hs = doc.styles[style_name]
        except KeyError:
            continue
        spec = heading_levels.get(lvl)
        if spec:
            # 仅覆盖规则指定的属性，保留 Word 默认标题层级感
            _set_style_font(hs, spec.name, spec.size_pt, spec.bold)
        # 标题不首行缩进（通用排版习惯）
        hs.paragraph_format.first_line_indent = Cm(0)

    # ── 3. 写入内容 ──
    if title:
        # Title 样式（审计时按主标题处理，不占标题层级、不按正文检查）
        p = doc.add_paragraph(title, style="Title")
        p.alignment = 1  # center
        if rules.body_font and rules.body_font.name:
            _set_style_font(doc.styles["Title"], rules.body_font.name, None)

    for block in _parse_md_blocks(content_md):
        btype = block["type"]
        if btype == "heading":
            doc.add_heading(block["text"], level=block["level"])
        elif btype == "list_item":
            doc.add_paragraph(block["text"], style="List Bullet")
        elif btype == "table":
            rows = block["rows"]
            if not rows:
                continue
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for i, row in enumerate(rows):
                for j, cell_text in enumerate(row):
                    if j < len(table.rows[i].cells):
                        table.rows[i].cells[j].text = cell_text
        else:  # para
            p = doc.add_paragraph(block["text"])
            if rules.body_first_line_indent_cm is not None:
                p.paragraph_format.first_line_indent = Cm(rules.body_first_line_indent_cm)

    # ── 4. 保存 + 自检 ──
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    logger.info(f"docx generated: {output_path.name}")

    self_check = audit_docx(output_path, rules)
    return {
        "status": "success",
        "file_path": output_path.name,
        "self_check": {
            "compliance_rate": self_check["compliance_rate"],
            "checks_passed": self_check["checks_passed"],
            "checks_total": self_check["checks_total"],
            "violations": self_check["violations"],
        },
    }
