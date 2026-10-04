"""pptx_fixer — PPT 轻量修复（V4.5）

只做显而易见的问题清理（不做美学排版）:
1. 删除空白页: 无任何 shapes，或所有 shape 均无内容（无文本/图片/表格/图表）
2. 清理空文本框: TEXT_BOX 类型且文本为空（保守策略: 不动占位符，不动含图片的 shape）

python-pptx 删除 slide 需操作 XML（sldIdLst + relationship）
"""
import logging
from pathlib import Path
from typing import Dict, Any, List

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

logger = logging.getLogger(__name__)


class PptxFixError(Exception):
    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.status_code = status_code


def _shape_has_content(shape) -> bool:
    """shape 是否携带实际内容"""
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return True
    if shape.has_table:
        return True
    if getattr(shape, "has_chart", False):
        return True
    if shape.has_text_frame and shape.text_frame.text.strip():
        return True
    return False


def _is_blank_slide(slide) -> bool:
    """整页无任何实际内容"""
    return not any(_shape_has_content(shape) for shape in slide.shapes)


def _delete_slide(prs: Presentation, slide_index: int) -> None:
    """从 Presentation 中删除指定索引的 slide（XML 级操作）"""
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    slide_id = slides[slide_index]
    # 删除 relationship
    rId = slide_id.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
    if rId:
        try:
            prs.part.drop_rel(rId)
        except Exception:
            pass
    xml_slides.remove(slide_id)


def fix_pptx(pptx_path: Path) -> Dict[str, Any]:
    """清理空白页与空文本框（写前自动 .bak 备份）"""
    if pptx_path.suffix.lower() != ".pptx":
        raise PptxFixError(f"仅支持 .pptx（当前: {pptx_path.suffix}）", 400)

    try:
        prs = Presentation(str(pptx_path))
    except Exception as e:
        raise PptxFixError(f"无法打开演示文稿（可能为 .ppt 老格式或已损坏）: {e}", 400)

    total_before = len(prs.slides)

    # 1. 清理每页的空文本框（不动占位符/图片/表格）
    removed_textboxes = 0
    for slide in prs.slides:
        to_remove = []
        for shape in slide.shapes:
            if (shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX
                    and shape.has_text_frame
                    and not shape.text_frame.text.strip()):
                to_remove.append(shape)
        for shape in to_remove:
            shape._element.getparent().remove(shape._element)
            removed_textboxes += 1

    # 2. 删除空白页（倒序删除避免索引位移）
    removed_slides: List[int] = []
    for i in range(len(prs.slides) - 1, -1, -1):
        if _is_blank_slide(prs.slides[i]):
            _delete_slide(prs, i)
            removed_slides.append(i)
    removed_slides.reverse()  # 正序展示

    # 备份 + 保存
    import shutil
    bak = pptx_path.with_suffix(pptx_path.suffix + ".bak")
    shutil.copy2(pptx_path, bak)
    prs.save(str(pptx_path))

    total_after = len(prs.slides)
    logger.info(f"pptx fixed: {pptx_path.name}, slides {total_before}->{total_after}, "
                f"textboxes removed={removed_textboxes}")

    return {
        "status": "success",
        "file_path": pptx_path.name,
        "slides_before": total_before,
        "slides_after": total_after,
        "removed_blank_slides": removed_slides,
        "removed_empty_textboxes": removed_textboxes,
        "backup": bak.name,
    }
