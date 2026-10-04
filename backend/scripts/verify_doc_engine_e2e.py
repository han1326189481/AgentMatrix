"""verify_doc_engine_e2e — Document Engine 端到端验证（V4.5）

验证四场景（零 LLM，纯组件级）:
  A. 夹具构造: 违规 docx + 规则 md + 带空白页 pptx
  B. detector: repair_docx / generate_docx / fix_pptx / read_only 判定
  C. 引擎执行: repair_docx / generate_docx / fix_pptx 落盘 + 复检
  D. Knowledge→Result 数据链路: doc_task 序列化/反序列化兼容性
"""
import sys
import json
import shutil
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_LINE_SPACING
from pptx import Presentation
from pptx.util import Inches

from core.document_engine.file_ops import get_workspace_dir, write_file
from core.document_engine.doc_task_detector import detect_doc_task
from core.document_engine.format_rules import from_rules_doc, FormatRules
from core.document_engine.pipeline import repair_docx
from core.document_engine.docx_generator import generate_docx
from core.document_engine.docx_audit import audit_docx
from core.document_engine.pptx_fixer import fix_pptx
from core.document_engine.docx_reader import parse_docx

WS = get_workspace_dir()
PASS, FAIL = 0, 0


def check(name: str, cond: bool, detail: str = ""):
    global PASS, FAIL
    mark = "PASS" if cond else "FAIL"
    print(f"  [{mark}] {name}" + (f" — {detail}" if detail else ""))
    if cond:
        PASS += 1
    else:
        FAIL += 1


def make_fixture_docx(path: Path):
    """构造违规 docx: 正文宋体五号(10.5pt)无首行缩进 + 标题黑体错误字号"""
    doc = Document()
    # 标题1: 用了黑体但字号错误（规则要求黑体三号16pt，这里给 14pt）
    h = doc.add_heading("项目背景", level=1)
    for run in h.runs:
        run.font.name = "黑体"
        run.font.size = Pt(14)
    # 正文: 宋体五号 10.5pt（规则要求仿宋小四12pt）
    p = doc.add_paragraph("本项目旨在构建一个多智能体协作系统，实现文档的自动化处理与格式规范化。")
    for run in p.runs:
        run.font.name = "宋体"
        run.font.size = Pt(10.5)
    # 无首行缩进、单倍行距（规则要求 1.5 倍行距 + 首行缩进 2 字符）
    p2 = doc.add_paragraph("系统采用本地模型与云端模型混合调度的策略，在保证响应速度的同时兼顾生成质量。")
    for run in p2.runs:
        run.font.name = "宋体"
        run.font.size = Pt(10.5)
    doc.save(str(path))


def make_fixture_pptx(path: Path):
    """构造问题 pptx: 2 空白页 + 1 含空文本框页 + 1 正常页"""
    prs = Presentation()
    # slide1: 空白页（无任何 shape）
    prs.slides.add_slide(prs.slide_layouts[6])
    # slide2: 正常内容页
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "系统架构介绍"
    # slide3: 只有空文本框的页
    s3 = prs.slides.add_slide(prs.slide_layouts[6])
    s3.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    # slide4: 空白页
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(str(path))


RULES_MD = """# 文档格式规范

## 页面设置
- 上边距 2.54cm，下边距 2.54cm，左边距 3.18cm，右边距 3.18cm

## 正文格式
- 正文字体：仿宋，字号小四
- 行距：1.5 倍行距
- 首行缩进 2 字符

## 标题格式
- 一级标题：黑体，三号
- 二级标题：黑体，四号
"""


def main():
    print("=" * 60)
    print("A. 夹具构造")
    print("=" * 60)
    fx = WS / "_e2e_fixture"
    if fx.exists():
        shutil.rmtree(fx)
    fx.mkdir(parents=True)
    docx_path = fx / "测试报告.docx"
    pptx_path = fx / "演示文稿.pptx"
    make_fixture_docx(docx_path)
    make_fixture_pptx(pptx_path)
    rules_rel = "_e2e_fixture/格式规范.md"
    write_file(rules_rel, RULES_MD)
    check("夹具创建", docx_path.exists() and pptx_path.exists() and (WS / rules_rel).exists())
    # 初始违规确认（修复前 audit 应低于 100%）
    pre = audit_docx(docx_path, from_rules_doc(WS / rules_rel))
    check("初始文档存在违规", pre["compliance_rate"] < 1.0, f"初始合规率 {pre['compliance_rate']:.0%}")

    print()
    print("=" * 60)
    print("B. detector 四场景判定")
    print("=" * 60)
    t1 = detect_doc_task("把 测试报告.docx 按格式规范.md 修复一下格式")
    check("repair_docx 判定", t1.task_type == "repair_docx",
          f"{t1.task_type} | target={t1.target_path} rules={t1.rules_path} | {t1.reason}")

    t2 = detect_doc_task("按格式规范.md 生成一份项目方案文档")
    check("generate_docx 判定", t2.task_type == "generate_docx"
          and t2.output_path == "generated/项目方案文档.docx",
          f"{t2.task_type} | rules={t2.rules_path} out={t2.output_path} | {t2.reason}")

    t3 = detect_doc_task("清理 演示文稿.pptx 的空白页和空文本框")
    check("fix_pptx 判定", t3.task_type == "fix_pptx",
          f"{t3.task_type} | target={t3.target_path} | {t3.reason}")

    t4 = detect_doc_task("读一下测试报告.docx 帮我总结主要内容")
    check("read_only 判定", t4.task_type == "read_only",
          f"{t4.task_type} | target={t4.target_path} | {t4.reason}")

    t5 = detect_doc_task("今天天气怎么样，适合出去玩吗")
    check("非文档任务 → none", t5.task_type == "none")

    # 显式通道（context.documents）
    t6 = detect_doc_task("修复这个文档的格式", context={
        "documents": [
            {"path": "_e2e_fixture/测试报告.docx", "role": "target"},
            {"path": rules_rel, "role": "rules"},
        ]})
    check("显式通道 repair", t6.task_type == "repair_docx" and t6.detection_source == "explicit",
          f"{t6.task_type} | {t6.detection_source}")

    print()
    print("=" * 60)
    print("C. 引擎执行")
    print("=" * 60)
    # C1: repair_docx
    rep = repair_docx(docx_path, from_rules_doc(WS / rules_rel))
    rate = rep.get("final_audit", {}).get("compliance_rate", 0)
    check("repair_docx 达标", rep.get("status") == "compliant" and rate >= 0.9,
          f"rounds={rep.get('rounds')} 合规率={rate:.0%} status={rep.get('status')}")
    if rep.get("remaining_violations"):
        for v in rep["remaining_violations"][:5]:
            print(f"      - 未达标: {v['message']}")

    # C2: generate_docx（模拟 Writer 输出的 markdown）
    md = """# 智能文档系统方案

## 系统概述

本系统面向企业内部的文档自动化处理场景，提供格式修复、合规生成与批量清理三类核心能力。系统全部格式处理在本地完成，不上传任何文档内容到云端，保障数据隐私安全。

## 核心功能

格式修复采用审计-补丁回路机制，先对文档进行全量格式审计，再生成机械补丁逐项修正，多轮迭代直至合规。文档生成支持从模板抽取样式规则，确保输出与既有模板完全一致。
"""
    out_abs = fx / "generated" / "生成方案.docx"
    gen = generate_docx(md, from_rules_doc(WS / rules_rel), out_abs)
    sc = gen.get("self_check", {})
    check("generate_docx 生成+自检", out_abs.exists() and sc.get("compliance_rate", 0) >= 0.9,
          f"自检合规率={sc.get('compliance_rate', 0):.0%}")
    # 内容核验: 生成文档可被解析回 markdown
    parsed = parse_docx(out_abs)
    check("生成文档内容完整", "文档自动化" in parsed["content_md"] and "核心功能" in parsed["content_md"],
          f"内容 {len(parsed['content_md'])} 字")

    # C3: fix_pptx（slide3 删空文本框后变空白页也被删 → 4→1）
    fixed = fix_pptx(pptx_path)
    check("fix_pptx 空白页删除", fixed["slides_before"] == 4 and fixed["slides_after"] == 1,
          f"{fixed['slides_before']} → {fixed['slides_after']} 页，"
          f"删空白页 {len(fixed.get('removed_blank_slides', []))}，"
          f"清空文本框 {fixed.get('removed_empty_textboxes', 0)} 个")

    print()
    print("=" * 60)
    print("D. Knowledge→Result doc_task 序列化链路")
    print("=" * 60)
    # 模拟 Knowledge Agent 序列化 → Result Agent 反序列化
    task_dict = t1.to_dict()
    serialized = json.dumps({"doc_task": task_dict}, ensure_ascii=False)
    restored = json.loads(serialized).get("doc_task", {})
    check("doc_task round-trip", restored.get("task_type") == "repair_docx"
          and restored.get("target_path") == t1.target_path,
          f"keys={sorted(restored.keys())}")

    print()
    print("=" * 60)
    print(f"结果: {PASS} PASS / {FAIL} FAIL")
    print("=" * 60)
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
